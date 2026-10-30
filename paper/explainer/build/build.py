"""Typeset content.py as Word (.docx) and PDF.

The .docx is written with python-docx. The PDF is typeset from the same blocks as HTML and printed by Chrome,
because Word's scripted PDF export replaces Apple SD Gothic Neo with substitute fonts.
Run:  python3 paper/explainer/build/build.py            (after make_figs.py)
Outputs: paper/explainer/out/Near-Miss_Blindness_연구설명.{docx,html,pdf}
"""
import datetime
import html
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from common import FIG, FONT_FAMILY, OUT
from content import AUTHOR, DOC, SUBTITLE, TITLE

BASENAME = "Near-Miss_Blindness_연구설명"
TEXT_W_CM = 17.0  # A4 width 21 cm minus 2 cm margins on each side

INK, INK2, NAVY = "2E2D2A", "52514E", "1F3A5F"
HEAD_FILL, CALLOUT_FILL, FORMULA_FILL = "EEF3F9", "F3F7FC", "F6F6F3"
RULE, HAIRLINE, ACCENT = "B5B3AD", "E4E3DF", "2A78D6"


# ---------------------------------------------------------------- shared: numbering and inline markup

def number_refs(doc):
    """Number tables and figures in order of appearance and resolve [[표:key]] / [[그림:key]]."""
    nums, t, f = {}, 0, 0
    for kind, b in doc:
        if kind == "table":
            t += 1
            nums[("표", b["key"])] = t
        elif kind == "figure":
            f += 1
            nums[("그림", b["key"])] = f

    def fix(s):
        def rep(m):
            k = (m.group(1), m.group(2))
            if k not in nums:
                raise KeyError(f"unknown reference {m.group(0)}")
            return f"{k[0]} {nums[k]}"
        return re.sub(r"\[\[(표|그림):([a-z0-9_]+)\]\]", rep, s)

    out = []
    for kind, b in doc:
        if kind in ("p", "h1", "h2"):
            out.append((kind, fix(b)))
        elif kind in ("bullets", "numbers"):
            out.append((kind, [fix(x) for x in b]))
        elif kind == "table":
            b = dict(b, caption=fix(b["caption"]), rows=[[fix(c) for c in r] for r in b["rows"]],
                     note=fix(b["note"]) if b["note"] else None, num=nums[("표", b["key"])])
            out.append((kind, b))
        elif kind == "figure":
            out.append((kind, dict(b, caption=fix(b["caption"]), num=nums[("그림", b["key"])])))
        elif kind == "callout":
            out.append((kind, dict(b, text=fix(b["text"]))))
        else:
            out.append((kind, b))
    return out


_INLINE = re.compile(r"\*\*(.+?)\*\*|_\{([^}]*)\}|\^\{([^}]*)\}")


def parse(text, bold=False):
    """Split mini-markup into (text, bold, vert) runs; vert is '', 'sub' or 'sup'."""
    runs, pos = [], 0
    for m in _INLINE.finditer(text):
        if m.start() > pos:
            runs.append((text[pos:m.start()], bold, ""))
        if m.group(1) is not None:
            runs.extend(parse(m.group(1), bold=True))
        elif m.group(2) is not None:
            runs.append((m.group(2), bold, "sub"))
        else:
            runs.append((m.group(3), bold, "sup"))
        pos = m.end()
    if pos < len(text):
        runs.append((text[pos:], bold, ""))
    return runs


# ---------------------------------------------------------------- DOCX

def _set_fonts(rpr):
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in list(rfonts.attrib):
        del rfonts.attrib[attr]
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), FONT_FAMILY)


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _borders(table, spec):
    """spec: {edge: (val, size_eighths, color)} for top/left/bottom/right/insideH/insideV."""
    tblpr = table._tbl.tblPr
    old = tblpr.find(qn("w:tblBorders"))
    if old is not None:
        tblpr.remove(old)
    b = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        val, sz, color = spec.get(edge, ("nil", 0, "auto"))
        e.set(qn("w:val"), val)
        if val != "nil":
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), color)
        b.append(e)
    tblpr.append(b)


def _cell_margins(table, top=50, left=90, bottom=50, right=90):
    tblpr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for edge, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:w"), str(v))
        e.set(qn("w:type"), "dxa")
        mar.append(e)
    tblpr.append(mar)


def _row_flag(row, tag):
    trpr = row._tr.get_or_add_trPr()
    trpr.append(OxmlElement(f"w:{tag}"))


def _add_runs(par, text, size=None, color=None, bold=False):
    for i, line in enumerate(text.split("\n")):
        if i:
            par.add_run().add_break()
        for seg, b, vert in parse(line, bold=bold):
            r = par.add_run(seg)
            r.bold = b or None
            if vert == "sub":
                r.font.subscript = True
            elif vert == "sup":
                r.font.superscript = True
            if size:
                r.font.size = Pt(size)
            if color:
                r.font.color.rgb = RGBColor.from_string(color)


def _fmt(par, before=0, after=6, line=1.42, align=None, keep=False):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if align is not None:
        par.alignment = align
    if keep:
        pf.keep_with_next = True


def _field(par, instr, size, color):
    for kind, text in (("begin", None), ("instr", instr), ("separate", None), ("text", "1"), ("end", None)):
        r = par.add_run()
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor.from_string(color)
        if kind == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = text
            r._r.append(it)
        elif kind == "text":
            r.text = text
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), kind)
            r._r.append(fc)


def _setup_docx():
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(1.9), Cm(1.9)
    sec.header_distance, sec.footer_distance = Cm(0.9), Cm(0.9)

    styles = d.styles.element
    rpr_default = styles.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    _set_fonts(rpr_default)
    lang = rpr_default.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr_default.append(lang)
    lang.set(qn("w:val"), "ko-KR")
    lang.set(qn("w:eastAsia"), "ko-KR")
    for st in styles.findall(qn("w:style")):
        rpr = st.find(qn("w:rPr"))
        if rpr is not None and rpr.find(qn("w:rFonts")) is not None:
            _set_fonts(rpr)

    normal = d.styles["Normal"]
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.42
    for name, size, before, after in (("Heading 1", 15, 18, 8), ("Heading 2", 12, 12, 5)):
        st = d.styles[name]
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.italic = False
        st.font.color.rgb = RGBColor.from_string(NAVY)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.2
        st.paragraph_format.keep_with_next = True

    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(foot, "PAGE", 8.5, "8A8883")
    r = foot.add_run(" / ")
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string("8A8883")
    _field(foot, "NUMPAGES", 8.5, "8A8883")

    cp = d.core_properties
    cp.title, cp.author, cp.subject = f"{TITLE}: {SUBTITLE}", AUTHOR, "연구설명"
    cp.created = datetime.datetime.now()
    return d


def _docx_table(d, b):
    cap = d.add_paragraph()
    _fmt(cap, before=8, after=3, line=1.2, keep=True)
    _add_runs(cap, f"표 {b['num']}. {b['caption']}", size=9, bold=True)

    header, rows = b["header"], b["rows"]
    ncol = len(header)
    table = d.add_table(rows=1 + len(rows), cols=ncol)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    tblpr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tblpr.append(layout)
    _borders(table, {"top": ("single", 8, RULE), "bottom": ("single", 8, RULE), "insideH": ("single", 4, HAIRLINE)})
    _cell_margins(table)
    widths = [Cm(TEXT_W_CM * w / sum(b["widths"])) for w in b["widths"]]
    for i, col in enumerate(table.columns):
        col.width = widths[i]
    fs = b["font"] or 8.6
    align = {"l": WD_ALIGN_PARAGRAPH.LEFT, "c": WD_ALIGN_PARAGRAPH.CENTER, "r": WD_ALIGN_PARAGRAPH.RIGHT}
    keep_whole = len(rows) <= 9
    for ri, values in enumerate([header] + rows):
        row = table.rows[ri]
        _row_flag(row, "cantSplit")
        if ri == 0:
            _row_flag(row, "tblHeader")
        for ci, text in enumerate(values):
            cell = row.cells[ci]
            cell.width = widths[ci]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            par = cell.paragraphs[0]
            _fmt(par, after=0, line=1.18, align=WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else align[b["align"][ci]],
                 keep=keep_whole and ri < len(rows))
            if ri == 0:
                _shade(cell, HEAD_FILL)
                _add_runs(par, text, size=fs, bold=True, color="1F2D3D")
            else:
                _add_runs(par, text, size=fs, bold=(ri - 1) in b["bold_rows"])
    after = d.add_paragraph()
    if b["note"]:
        _fmt(after, before=3, after=8, line=1.25)
        _add_runs(after, b["note"], size=8.2, color=INK2)
    else:
        _fmt(after, after=4, line=1.0)
        after.add_run().font.size = Pt(4)


def _docx_box(d, lines, fill, accent=None, center=False, size=None):
    table = d.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    spec = {"left": ("single", 18, accent)} if accent else {}
    _borders(table, spec)
    _cell_margins(table, top=110, left=170, bottom=110, right=170)
    cell = table.rows[0].cells[0]
    cell.width = Cm(TEXT_W_CM)
    table.columns[0].width = Cm(TEXT_W_CM)
    _shade(cell, fill)
    for i, line in enumerate(lines):
        par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        _fmt(par, after=2 if i < len(lines) - 1 else 0, line=1.4,
             align=WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY)
        _add_runs(par, line, size=size)
    spacer = d.add_paragraph()
    _fmt(spacer, after=4, line=1.0)
    spacer.add_run().font.size = Pt(4)


def build_docx(doc, path):
    d = _setup_docx()
    p = d.add_paragraph()
    _fmt(p, before=6, after=4, line=1.2, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_runs(p, TITLE, size=21, bold=True, color="1F2F4A")
    p = d.add_paragraph()
    _fmt(p, after=8, line=1.25, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_runs(p, SUBTITLE, size=12.5, color="4A5A72")
    p = d.add_paragraph()
    _fmt(p, after=10, line=1.2, align=WD_ALIGN_PARAGRAPH.RIGHT)
    _add_runs(p, AUTHOR, size=10.5, color=INK2)

    for kind, b in doc:
        if kind == "h1":
            h = d.add_heading(level=1)
            _add_runs(h, b)
            ppr = h._p.get_or_add_pPr()
            bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "3"), ("w:color", "CDD6E3")):
                bottom.set(qn(k), v)
            bdr.append(bottom)
            ppr.append(bdr)
        elif kind == "h2":
            h = d.add_heading(level=2)
            _add_runs(h, b)
        elif kind == "p":
            p = d.add_paragraph()
            label = b.startswith("**") and b.endswith("**")  # a bold one-line label stays with what follows
            _fmt(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep=label)
            _add_runs(p, b)
        elif kind in ("bullets", "numbers"):
            for item in b:
                p = d.add_paragraph(style="List Bullet" if kind == "bullets" else "List Number")
                _fmt(p, after=3.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
                _add_runs(p, item)
            d.paragraphs[-1].paragraph_format.space_after = Pt(8)
        elif kind == "table":
            _docx_table(d, b)
        elif kind == "figure":
            p = d.add_paragraph()
            _fmt(p, before=4, after=2, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER, keep=True)
            p.add_run().add_picture(os.path.join(FIG, b["file"]), width=Cm(16.6 * b["width"]))
            cap = d.add_paragraph()
            _fmt(cap, after=10, line=1.3, align=WD_ALIGN_PARAGRAPH.CENTER)
            _add_runs(cap, f"그림 {b['num']}. {b['caption']}", size=8.6, color=INK2)
        elif kind == "callout":
            text = (f"**{b['title']}.** " if b["title"] else "") + b["text"]
            _docx_box(d, [text], CALLOUT_FILL, accent=ACCENT)
        elif kind == "formula":
            _docx_box(d, b, FORMULA_FILL, center=True, size=9.6)
        elif kind == "pagebreak":
            d.add_page_break()
    d.save(path)
    return path


# ---------------------------------------------------------------- HTML -> PDF

CSS = """
@page { size: A4; margin: 19mm 20mm 21mm 20mm; }
html { -webkit-print-color-adjust: exact; print-color-adjust: exact; }
body { font-family: "Apple SD Gothic Neo", sans-serif; font-size: 10pt; line-height: 1.74; color: #2e2d2a;
       word-break: keep-all; overflow-wrap: break-word; margin: 0; font-kerning: normal; }
.title { font-size: 21pt; font-weight: 700; color: #1f2f4a; text-align: center; margin: 2pt 0 3pt; letter-spacing: -0.2pt; }
.subtitle { font-size: 12.5pt; font-weight: 500; color: #4a5a72; text-align: center; margin: 0 0 8pt; }
.author { text-align: right; font-size: 10.5pt; color: #52514e; margin: 0 0 6pt; }
h1 { font-size: 15pt; font-weight: 700; color: #1f3a5f; border-bottom: 1px solid #cdd6e3; padding-bottom: 4pt;
     margin: 20pt 0 10pt; line-height: 1.35; break-after: avoid; }
h2 { font-size: 12pt; font-weight: 600; color: #1f3a5f; margin: 15pt 0 5pt; line-height: 1.4; break-after: avoid; }
p { margin: 0 0 7pt; text-align: justify; }
p.label { margin: 10pt 0 4pt; break-after: avoid; }
ul, ol { margin: 0 0 9pt; padding-left: 15pt; }
li { margin: 0 0 4pt; text-align: justify; padding-left: 2pt; }
li::marker { color: #6f8fb8; }
b, strong { font-weight: 600; color: #1f1e1c; }
sub, sup { font-size: 68%; line-height: 0; }
.tcap { font-size: 9pt; font-weight: 600; color: #2e2d2a; margin: 10pt 0 4pt; break-after: avoid; }
table { border-collapse: collapse; width: 100%; font-size: 8.6pt; line-height: 1.5; margin: 0 0 4pt;
        border-top: 1.1px solid #b5b3ad; border-bottom: 1.1px solid #b5b3ad; }
table.keep { break-inside: avoid; }
th { background: #eef3f9; font-weight: 600; color: #1f2d3d; padding: 4pt 5pt; text-align: center; vertical-align: middle;
     border-bottom: 0.8px solid #c9d3e0; }
td { padding: 3.6pt 5pt; border-top: 0.6px solid #e4e3df; vertical-align: middle; }
tr { break-inside: avoid; }
td.l { text-align: left; } td.c { text-align: center; } td.r { text-align: right; }
tr.em td { font-weight: 600; color: #1f1e1c; }
.tnote { font-size: 8.2pt; color: #52514e; margin: 3pt 0 10pt; line-height: 1.5; }
.tgap { height: 8pt; }
.tblock { break-inside: avoid; }
figure { margin: 8pt 0 12pt; text-align: center; break-inside: avoid; }
figure img { display: block; margin: 0 auto; }
figcaption { font-size: 8.6pt; color: #52514e; margin-top: 5pt; line-height: 1.55; text-align: center; }
.callout { background: #f3f7fc; border-left: 3px solid #2a78d6; border-radius: 3px; padding: 7pt 11pt;
           margin: 4pt 0 12pt; break-inside: avoid; text-align: justify; }
.formula { background: #f6f6f3; border-radius: 4px; padding: 7pt 12pt; margin: 3pt 0 11pt; text-align: center;
           font-size: 9.6pt; line-height: 1.7; break-inside: avoid; }
.pagebreak { break-before: page; }
"""

FOOTER = ('<div style="width:100%;text-align:center;font-family:\'Apple SD Gothic Neo\';font-size:8pt;'
          'color:#8a8883;"><span class="pageNumber"></span> / <span class="totalPages"></span></div>')


def _h(text):
    out = []
    for i, line in enumerate(text.split("\n")):
        if i:
            out.append("<br>")
        for seg, bold, vert in parse(line):
            s = html.escape(seg)
            if vert:
                s = f"<{vert}>{s}</{vert}>"
            if bold:
                s = f"<b>{s}</b>"
            out.append(s)
    return "".join(out)


def build_html(doc, path):
    parts = [f'<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>{html.escape(TITLE)}</title>'
             f"<style>{CSS}</style></head><body>",
             f'<div class="title">{_h(TITLE)}</div><div class="subtitle">{_h(SUBTITLE)}</div>'
             f'<div class="author">{_h(AUTHOR)}</div>']
    rel = os.path.relpath(FIG, os.path.dirname(path))
    for kind, b in doc:
        if kind == "h1":
            parts.append(f"<h1>{_h(b)}</h1>")
        elif kind == "h2":
            parts.append(f"<h2>{_h(b)}</h2>")
        elif kind == "p":
            label = b.startswith("**") and b.endswith("**")
            parts.append(f'<p class="label">{_h(b)}</p>' if label else f"<p>{_h(b)}</p>")
        elif kind in ("bullets", "numbers"):
            tag = "ul" if kind == "bullets" else "ol"
            parts.append(f"<{tag}>" + "".join(f"<li>{_h(x)}</li>" for x in b) + f"</{tag}>")
        elif kind == "table":
            tot = sum(b["widths"])
            cols = "".join(f'<col style="width:{100 * w / tot:.2f}%">' for w in b["widths"])
            fs = f' style="font-size:{b["font"]}pt"' if b["font"] else ""
            keep = len(b["rows"]) <= 9
            head = "".join(f"<th>{_h(x)}</th>" for x in b["header"])
            body = "".join(
                f'<tr class="{"em" if i in b["bold_rows"] else ""}">'
                + "".join(f'<td class="{b["align"][j]}">{_h(x)}</td>' for j, x in enumerate(r)) + "</tr>"
                for i, r in enumerate(b["rows"]))
            note = f'<div class="tnote">{_h(b["note"])}</div>' if b["note"] else '<div class="tgap"></div>'
            block = (f'<div class="tcap">표 {b["num"]}. {_h(b["caption"])}</div>'
                     f'<table class="{"keep" if keep else ""}"{fs}><colgroup>{cols}</colgroup>'
                     f"<thead><tr>{head}</tr></thead><tbody>{body}</tbody></table>{note}")
            parts.append(f'<div class="tblock">{block}</div>' if keep else block)
        elif kind == "figure":
            parts.append(f'<figure><img src="{rel}/{b["file"]}" style="width:{100 * b["width"]:.1f}%">'
                         f'<figcaption>그림 {b["num"]}. {_h(b["caption"])}</figcaption></figure>')
        elif kind == "callout":
            title = f"<b>{_h(b['title'])}.</b> " if b["title"] else ""
            parts.append(f'<div class="callout">{title}{_h(b["text"])}</div>')
        elif kind == "formula":
            parts.append('<div class="formula">' + "<br>".join(_h(x) for x in b) + "</div>")
        elif kind == "pagebreak":
            parts.append('<div class="pagebreak"></div>')
    parts.append("</body></html>")
    with open(path, "w", encoding="utf-8") as f:
        f.write("".join(parts))
    return path


def build_pdf(html_path, pdf_path):
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch(channel="chrome")
        page = browser.new_page()
        page.goto("file://" + html_path)
        page.wait_for_load_state("networkidle")
        page.evaluate("document.fonts.ready")
        page.pdf(path=pdf_path, format="A4", print_background=True, display_header_footer=True,
                 header_template="<span></span>", footer_template=FOOTER,
                 margin={"top": "19mm", "bottom": "21mm", "left": "20mm", "right": "20mm"})
        browser.close()
    return pdf_path


def main(argv):
    os.makedirs(OUT, exist_ok=True)
    doc = number_refs(DOC)
    targets = set(argv) or {"docx", "pdf"}
    if "docx" in targets:
        print("wrote", build_docx(doc, os.path.join(OUT, BASENAME + ".docx")))
    if "pdf" in targets:
        h = build_html(doc, os.path.join(OUT, BASENAME + ".html"))
        print("wrote", h)
        print("wrote", build_pdf(h, os.path.join(OUT, BASENAME + ".pdf")))


if __name__ == "__main__":
    main(sys.argv[1:])
